"""
Generate a Menoufia University / Faculty of Artificial Intelligence graduation book.

This script does not fill an existing DOCX. It uses the uploaded DOCX as a
blueprint and rebuilds the same document structure in Python:

- English cover page with university/faculty logos
- front matter: acknowledgment, abstract, table of contents placeholder,
  lists of tables/figures, abbreviations, project summary
- numbered chapters with headers, footers, section headings, paragraphs,
  figures, captions and tables
- bibliography
- Arabic summary and Arabic cover page

Install:
    pip install -r requirements.txt

Run:
    python generate_graduation_book.py

Output:
    generated_graduation_book.docx

Content:
    Edit content/book.json and content/chapters/*.md.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
ASSET_DIR = BASE_DIR / "assets"
CONTENT_DIR = BASE_DIR / "content"
BOOK_JSON = CONTENT_DIR / "book.json"

OUTPUT_DOCX = BASE_DIR / "generated_graduation_book.docx"

# Extracted from the uploaded DOCX.
MENOUFIA_LOGO = ASSET_DIR / "menoufia_logo.jpeg"
FACULTY_AI_LOGO = ASSET_DIR / "faculty_ai_logo.png"

TEXT_WIDTH_INCHES = 6.2
LIGHT_GRID = "C8D1CD"
MID_GRID = "9EAAA5"
HEADER_FILL = "F2F5F4"
ACCENT_TEXT = RGBColor(31, 49, 47)
MUTED_TEXT = RGBColor(82, 99, 95)
LTR_TOKEN_RE = re.compile(r"(?<![\w-])([A-Za-z][A-Za-z0-9.+/-]*)(?![\w-])")
PAGE_BREAK_MARKER = "[[PAGE_BREAK]]"

NUMERIC_BOOKTABS_TABLES = {
    "Table 2.4",
    "Table 6.3",
    "Table 7.5",
    "Table 9.5",
    "Table 9.6",
    "Table 9.7",
    "Table 9.9",
}

DENSE_GRID_TABLES = {
    "Table 2.2",
    "Table 2.3",
    "Table 4.1",
    "Table 4.2",
    "Table 4.3",
    "Table 4.4",
    "Table 4.5",
    "Table 5.2",
    "Table 6.4",
    "Table 9.8",
    "Table 9.10",
    "Table 10.2",
}


# -----------------------------------------------------------------------------
# Low-level helpers
# -----------------------------------------------------------------------------

def set_cell_text(cell, text: str, *, bold: bool = False, size: int | float | None = None,
                  align: int | None = None, rtl: bool = False,
                  color: RGBColor | None = None, line_spacing: float = 1.05):
    """Replace a cell's first paragraph with styled text."""
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = line_spacing
    if rtl:
        set_rtl(p)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, *, size: int | float | None = None, bold: bool | None = None,
                 italic: bool | None = None, color: RGBColor | None = None,
                 font: str = "Times New Roman"):
    """Set Latin/complex-script fonts reliably."""
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_run_rtl(run, rtl: bool):
    """Set run-level text direction explicitly for mixed Arabic/Latin prose."""
    rpr = run._element.get_or_add_rPr()
    rtl_el = rpr.find(qn("w:rtl"))
    if rtl_el is None:
        rtl_el = OxmlElement("w:rtl")
        rpr.append(rtl_el)
    rtl_el.set(qn("w:val"), "1" if rtl else "0")


def set_run_lang(run, *, val: str | None = None, bidi: str | None = None):
    """Set Word language hints so Latin tokens inside RTL paragraphs stay LTR."""
    rpr = run._element.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    if val:
        lang.set(qn("w:val"), val)
    if bidi:
        lang.set(qn("w:bidi"), bidi)


def add_text(paragraph, text: str, *, size: int | float | None = None,
             bold: bool = False, italic: bool = False,
             color: RGBColor | None = None, font: str = "Times New Roman",
             rtl: bool | None = None, lang: str | None = None):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color, font=font)
    if rtl is not None:
        set_run_rtl(run, rtl)
    if lang is not None:
        set_run_lang(run, val=lang, bidi="ar-SA" if rtl else None)
    return run


def iter_ltr_token_parts(text: str):
    """Yield (part, is_ltr) pairs for Latin tokens inside Arabic prose."""
    pos = 0
    for match in LTR_TOKEN_RE.finditer(text):
        if match.start() > pos:
            yield text[pos:match.start()], False
        yield match.group(1), True
        pos = match.end()
    if pos < len(text):
        yield text[pos:], False


def add_mixed_rtl_text(paragraph, text: str, *, size: int | float = 14,
                       bold: bool = False):
    """Add Arabic prose with embedded Latin/acronym tokens as real LTR runs."""
    for part, is_ltr in iter_ltr_token_parts(text):
        if not part:
            continue
        if is_ltr:
            add_text(paragraph, part, size=size, bold=bold, rtl=False, lang="en-US")
        else:
            add_text(paragraph, part, size=size, bold=bold, rtl=True, lang="ar-SA")


def set_rtl(paragraph):
    """Set right-to-left paragraph direction in WordprocessingML."""
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_table_borders(table, color="000000", size="4"):
    """Apply visible borders to a table."""
    set_table_border_edges(
        table,
        {
            "top": ("single", size, color),
            "left": ("single", size, color),
            "bottom": ("single", size, color),
            "right": ("single", size, color),
            "insideH": ("single", size, color),
            "insideV": ("single", size, color),
        },
    )


def set_table_border_edges(table, edges: dict[str, tuple[str, str, str]]):
    """Apply explicit border settings by edge name."""
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        val, size, color = edges.get(edge, ("nil", "0", "auto"))
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width: Cm):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_width_inches(cell, width: float):
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, *, top: int = 80, start: int = 100,
                     bottom: int = 80, end: int = 100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_row_border(row, *, top: tuple[str, str, str] | None = None,
                   bottom: tuple[str, str, str] | None = None):
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        for edge, spec in {"top": top, "bottom": bottom}.items():
            if spec is None:
                continue
            tag = f"w:{edge}"
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            val, size, color = spec
            element.set(qn("w:val"), val)
            element.set(qn("w:sz"), size)
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), color)


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def keep_table_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "1")


def clear_container(container):
    """Remove existing paragraphs/tables from a header/footer/body-like container."""
    element = container._element
    for child in list(element):
        element.remove(child)


def add_page_number(paragraph):
    """Insert a Word PAGE field into a paragraph."""
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    text = OxmlElement("w:t")
    text.text = "1"
    run_text = OxmlElement("w:r")
    run_text.append(text)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    for node in (fld_begin, instr, fld_sep, run_text, fld_end):
        run = OxmlElement("w:r") if node.tag != qn("w:r") else node
        if node.tag != qn("w:r"):
            run.append(node)
        paragraph._p.append(run)


def add_horizontal_border(paragraph, where="bottom"):
    """Add a single line above/below a paragraph."""
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    border = OxmlElement(f"w:{where}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), "000000")
    pbdr.append(border)


# -----------------------------------------------------------------------------
# Document setup
# -----------------------------------------------------------------------------

def configure_document(doc: Document):
    """A4 page, margins, and base styles matching the uploaded DOCX."""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.2)
        section.footer_distance = Cm(1.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.font.size = Pt(14)

    # The uploaded file relies heavily on Normal + manual formatting.
    # These style definitions make our generator easier to read.
    styles["Heading 1"].font.name = "Times New Roman"
    styles["Heading 1"].font.size = Pt(22)
    styles["Heading 1"].font.bold = True

    styles["Heading 2"].font.name = "Times New Roman"
    styles["Heading 2"].font.size = Pt(18)
    styles["Heading 2"].font.bold = True

    styles["Heading 3"].font.name = "Times New Roman"
    styles["Heading 3"].font.size = Pt(16)
    styles["Heading 3"].font.bold = True


def new_section(doc: Document, *, with_header: bool = False,
                chapter_no: int | None = None, chapter_title: str | None = None):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    clear_container(section.header)
    clear_container(section.footer)
    if with_header and chapter_no is not None and chapter_title:
        set_chapter_header_footer(section, chapter_no, chapter_title)
    return section


def set_chapter_header_footer(section, chapter_no: int, chapter_title: str):
    header = section.header
    p = header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, f"CHAPTER {chapter_no}", size=10, bold=False)
    p.add_run("\t")
    add_text(p, chapter_title.upper(), size=10, bold=True)
    add_horizontal_border(p, "bottom")

    footer = section.footer
    p = footer.add_paragraph()
    add_horizontal_border(p, "top")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "- ", size=12)
    add_page_number(p)
    add_text(p, " -", size=12)


# -----------------------------------------------------------------------------
# High-level blocks matching the uploaded DOCX
# -----------------------------------------------------------------------------

def add_cover_page(doc: Document, data: dict):
    """English cover page from page 1 of the uploaded document."""
    # The source cover page uses a tighter top/bottom layout than the body.
    doc.sections[0].top_margin = Cm(1.2)
    doc.sections[0].bottom_margin = Cm(1.0)
    doc.sections[0].left_margin = Cm(2.5)
    doc.sections[0].right_margin = Cm(2.5)
    # Header logo row: faculty logo | university/faculty text | university logo.
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table, color="FFFFFF", size="0")

    cells = table.rows[0].cells
    set_cell_width(cells[0], Cm(4.0))
    set_cell_width(cells[1], Cm(9.5))
    set_cell_width(cells[2], Cm(4.0))
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells[0].paragraphs[0].add_run().add_picture(str(FACULTY_AI_LOGO), width=Inches(0.9))

    p = cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = add_text(p, data["university"], size=14)
    run.font.small_caps = True
    p.add_run().add_break()
    run = add_text(p, data["faculty"], size=14)
    run.font.small_caps = True

    cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells[2].paragraphs[0].add_run().add_picture(str(MENOUFIA_LOGO), width=Inches(0.85))

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, data["project_title"], size=15, bold=True)

    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "A Graduation Project Submitted for the Degree of", size=14)
    p.add_run().add_break()
    add_text(p, "Bachelor in Artificial Intelligence", size=14)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Project Team:", size=16, bold=True, italic=True)

    add_team_table(doc, data["team"], rtl=False)

    for _ in range(1):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Supervised by:", size=18, bold=True)

    for _ in range(1):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for idx, line in enumerate(data["supervisor_lines"]):
        if idx:
            p.add_run().add_break()
        add_text(p, line, size=14)

    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, data["term"], size=16, bold=True)


def add_team_table(doc: Document, team: Sequence[dict], *, rtl: bool = False):
    rows = len(team) + 1
    table = doc.add_table(rows=rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)

    headers = ("Name", "Department") if not rtl else ("القسم العلمي", "الاسم")
    for idx, header in enumerate(headers):
        align = WD_ALIGN_PARAGRAPH.CENTER
        if not rtl and idx == 0:
            align = WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text(table.rows[0].cells[idx], header, bold=True, size=13,
                      align=align, rtl=rtl)

    for i, member in enumerate(team, start=1):
        dept = member.get("department", "")
        name = member.get("name", "")
        if rtl:
            set_cell_text(table.rows[i].cells[0], dept, bold=True, size=11,
                          align=WD_ALIGN_PARAGRAPH.CENTER, rtl=True)
            set_cell_text(table.rows[i].cells[1], name, bold=True, size=11,
                          align=WD_ALIGN_PARAGRAPH.CENTER, rtl=True)
        else:
            set_cell_text(table.rows[i].cells[0], name, size=12,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(table.rows[i].cells[1], dept, size=12,
                          align=WD_ALIGN_PARAGRAPH.CENTER)
    return table


def add_front_heading(doc: Document, title: str):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # First-letter-large look like the uploaded front-matter headings.
    if title:
        add_text(p, title[0], size=30, bold=True)
        add_text(p, title[1:].upper(), size=22, bold=True)
    return p


def add_front_matter(doc: Document, data: dict):
    new_section(doc)
    add_front_heading(doc, "ACKNOWLEDGMENT")
    for paragraph in data["acknowledgment"]:
        add_body_paragraph(doc, paragraph)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p, "Team Members", size=16, bold=True, font="Edwardian Script ITC")

    new_section(doc)
    add_front_heading(doc, "ABSTRACT")
    for paragraph in data["abstract"]:
        add_body_paragraph(doc, paragraph)

    new_section(doc)
    add_front_heading(doc, "TABLE OF CONTENTS")
    add_toc_placeholder(doc, data["toc"])

    new_section(doc)
    add_front_heading(doc, "LIST OF TABLES")
    for item in data.get("list_of_tables", []):
        add_toc_line(doc, item[0], item[1], level=1)

    new_section(doc)
    add_front_heading(doc, "LIST OF FIGURES")
    for item in data.get("list_of_figures", []):
        add_toc_line(doc, item[0], item[1], level=1)

    new_section(doc)
    add_front_heading(doc, "LIST OF ABBREVIATIONS")
    add_abbreviations_table(doc, data.get("abbreviations", []))

    new_section(doc)
    add_front_heading(doc, "PROJECT SUMMARY")
    for paragraph in data["project_summary"]:
        add_body_paragraph(doc, paragraph)


def add_toc_placeholder(doc: Document, toc_items: Sequence[dict]):
    """Static TOC, matching the uploaded DOCX enough for a generator skeleton.

    For a production thesis, use Word's References > Update Table after opening.
    """
    for item in toc_items:
        add_toc_line(doc, item["title"], str(item.get("page", "")), level=item.get("level", 1))


def add_toc_line(doc: Document, title: str, page: str, *, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5 * (level - 1))
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.2))
    add_text(p, title.upper() if level == 1 else title, size=12)
    add_text(p, "\t" + page, size=12)


def add_abbreviations_table(doc: Document, abbreviations: Sequence[tuple[str, str]]):
    table = doc.add_table(rows=max(2, len(abbreviations) + 1), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_cell_text(table.rows[0].cells[0], "Abbreviation", bold=True, size=12,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[0].cells[1], "Referenced Terms", bold=True, size=12,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, (abbr, meaning) in enumerate(abbreviations, start=1):
        set_cell_text(table.rows[i].cells[0], abbr, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.rows[i].cells[1], meaning, size=12, align=WD_ALIGN_PARAGRAPH.LEFT)


def add_chapter(doc: Document, chapter: dict):
    chapter_no = chapter["number"]
    chapter_title = chapter["title"]
    new_section(doc, with_header=True, chapter_no=chapter_no, chapter_title=chapter_title)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, f"Chapter {chapter_no}", size=24, bold=True, font="Arial")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, chapter_title.upper(), size=28, bold=True)

    for block in chapter.get("blocks", []):
        kind = block.get("type")
        if kind == "heading":
            add_section_heading(doc, block["text"], level=block.get("level", 2))
        elif kind == "paragraph":
            add_body_paragraph(doc, block["text"])
        elif kind == "page_break":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
        elif kind == "list":
            add_numbered_list(doc, block["items"])
        elif kind == "figure":
            add_figure(doc, ASSET_DIR / block["path"], block["caption"], width=block.get("width", 4.5))
        elif kind == "table":
            add_caption(doc, block["caption"], kind="table")
            add_data_table(doc, block["caption"], block["headers"], block["rows"])
        elif kind == "algorithm":
            add_algorithm_block(doc, block["title"], block["lines"])


def add_section_heading(doc: Document, text: str, *, level: int = 2):
    p = doc.add_paragraph()
    
    try:
        p.style = doc.styles[f"Heading {level}"]
    except KeyError:
        p.style = doc.styles["Heading 2"]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(p, text, size=18 if level == 2 else 16, bold=True)


def add_body_paragraph(doc: Document, text: str, *, rtl: bool = False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if rtl:
        set_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if rtl:
        add_mixed_rtl_text(p, text, size=14)
    else:
        add_text(p, text, size=14)
    return p


def add_numbered_list(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_text(p, item, size=14)


def add_caption(doc: Document, text: str, *, kind: str = "figure"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6 if kind == "table" else 4)
    p.paragraph_format.space_after = Pt(3 if kind == "table" else 6)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size=12, bold=True)


def add_figure(doc: Document, image_path: Path, caption: str, *, width: float = 4.5):
    if image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption, kind="figure")


def table_number(caption: str) -> str:
    return caption.split(":", 1)[0]


def looks_numeric(value: str) -> bool:
    text = value.strip().lower()
    if not text:
        return False
    return bool(re.search(r"(^|[\s/~=-])(?:about\s+)?\d", text)) or "r =" in text or "lambda" in text


def classify_table(caption: str, headers: Sequence[str], rows: Sequence[Sequence[str]]) -> str:
    number = table_number(caption)
    cells = [str(item) for item in headers]
    cells.extend(str(value) for row in rows for value in row)
    numeric_ratio = sum(1 for cell in cells if looks_numeric(cell)) / max(1, len(cells))
    long_ratio = sum(1 for cell in cells if len(cell) > 36) / max(1, len(cells))
    if number in NUMERIC_BOOKTABS_TABLES or numeric_ratio >= 0.45:
        return "compact_booktabs" if len(rows) >= 7 else "booktabs"
    if number in DENSE_GRID_TABLES or len(rows) >= 8 or (len(headers) >= 4 and long_ratio >= 0.35):
        return "compact_grid"
    return "clean_grid"


def column_numeric_flags(headers: Sequence[str], rows: Sequence[Sequence[str]]) -> list[bool]:
    flags = []
    for idx in range(len(headers)):
        values = [str(row[idx]) for row in rows if idx < len(row)]
        if not values:
            flags.append(False)
            continue
        flags.append(sum(1 for value in values if looks_numeric(value)) / len(values) >= 0.6)
    return flags


def column_widths(headers: Sequence[str], rows: Sequence[Sequence[str]], style: str) -> list[float]:
    if not headers:
        return []
    numeric_cols = column_numeric_flags(headers, rows)
    weights = []
    for idx, header in enumerate(headers):
        values = [str(header)] + [str(row[idx]) for row in rows if idx < len(row)]
        avg_len = sum(len(value) for value in values) / len(values)
        max_len = max(len(value) for value in values)
        if numeric_cols[idx]:
            weight = 0.9
        else:
            weight = 1.1 + min(max_len, 68) / 42 + min(avg_len, 42) / 55
        if idx == 0 and style.startswith("booktabs"):
            weight += 0.25
        weights.append(max(0.8, weight))
    total = sum(weights)
    return [TEXT_WIDTH_INCHES * weight / total for weight in weights]


def choose_cell_alignment(value: str, *, is_header: bool, numeric_column: bool,
                          style: str, is_first_column: bool = False) -> int:
    if is_header:
        return WD_ALIGN_PARAGRAPH.CENTER
    text = str(value).strip()
    if style.startswith("booktabs") and is_first_column and not numeric_column:
        return WD_ALIGN_PARAGRAPH.LEFT
    if style.startswith("booktabs") and numeric_column:
        if len(text) > 32 and text.count(" ") > 4:
            return WD_ALIGN_PARAGRAPH.LEFT
        return WD_ALIGN_PARAGRAPH.RIGHT
    if len(text) <= 20 and text.count(" ") <= 2:
        return WD_ALIGN_PARAGRAPH.CENTER
    return WD_ALIGN_PARAGRAPH.LEFT


def apply_table_shell(table, style: str):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    try:
        table.style = "Table Grid" if style.endswith("grid") else "Table Normal"
    except KeyError:
        pass
    if style.endswith("grid"):
        set_table_border_edges(
            table,
            {
                "top": ("single", "4", LIGHT_GRID),
                "left": ("single", "4", LIGHT_GRID),
                "bottom": ("single", "4", LIGHT_GRID),
                "right": ("single", "4", LIGHT_GRID),
                "insideH": ("single", "4", LIGHT_GRID),
                "insideV": ("single", "4", LIGHT_GRID),
            },
        )
    else:
        set_table_border_edges(
            table,
            {
                "top": ("single", "8", MID_GRID),
                "left": ("nil", "0", "auto"),
                "bottom": ("single", "8", MID_GRID),
                "right": ("nil", "0", "auto"),
                "insideH": ("nil", "0", "auto"),
                "insideV": ("nil", "0", "auto"),
            },
        )


def add_data_table(doc: Document, caption: str, headers: Sequence[str], rows: Sequence[Sequence[str]]):
    style = classify_table(caption, headers, rows)
    numeric_cols = column_numeric_flags(headers, rows)
    widths = column_widths(headers, rows, style)
    header_size = 10.7 if style == "compact_grid" else 11.2
    body_size = 10.3 if style in {"compact_grid", "compact_booktabs"} else 11.0
    margin = 68 if style in {"compact_grid", "compact_booktabs"} else 92
    line_spacing = 1.0 if style in {"compact_grid", "compact_booktabs"} else 1.08

    table = doc.add_table(rows=1, cols=len(headers))
    apply_table_shell(table, style)
    repeat_header_row(table.rows[0])
    keep_table_row_together(table.rows[0])

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width_inches(cell, widths[i])
        set_cell_margins(cell, top=margin, start=margin, bottom=margin, end=margin)
        if style.endswith("grid"):
            set_cell_shading(cell, HEADER_FILL)
        set_cell_text(
            cell,
            header,
            bold=True,
            size=header_size,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            color=ACCENT_TEXT,
            line_spacing=line_spacing,
        )
    if style.startswith("booktabs"):
        set_row_border(table.rows[0], bottom=("single", "6", MID_GRID))

    for row in rows:
        row_obj = table.add_row()
        keep_table_row_together(row_obj)
        cells = row_obj.cells
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_width_inches(cell, widths[i])
            set_cell_margins(cell, top=margin, start=margin, bottom=margin, end=margin)
            align = choose_cell_alignment(
                str(value),
                is_header=False,
                numeric_column=numeric_cols[i],
                style=style,
                is_first_column=i == 0,
            )
            set_cell_text(
                cell,
                str(value),
                size=body_size,
                align=align,
                color=MUTED_TEXT if align == WD_ALIGN_PARAGRAPH.LEFT else ACCENT_TEXT,
                line_spacing=line_spacing,
            )
    return table


def add_algorithm_block(doc: Document, title: str, lines: Sequence[str]):
    p = doc.add_paragraph()
    add_text(p, title, size=14, bold=True)
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        add_text(p, line, size=12)


def add_bibliography(doc: Document, refs: Sequence[str]):
    new_section(doc)
    add_front_heading(doc, "BIBLIOGRAPHY")
    for idx, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, f"[{idx}]. {ref}", size=12)


def add_arabic_summary(doc: Document, data: dict):
    new_section(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p)
    add_text(p, "مُلخـــــصُ المشروع", size=22, bold=True)
    for paragraph in data["arabic_summary"]:
        add_body_paragraph(doc, paragraph, rtl=True)

    new_section(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p)
    add_text(p, "المُوجَــــــــــزُ", size=22, bold=True)
    for paragraph in data["arabic_abstract"]:
        add_body_paragraph(doc, paragraph, rtl=True)


def add_arabic_cover_page(doc: Document, data: dict):
    section = new_section(doc)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(0.8)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table, color="FFFFFF", size="0")
    cells = table.rows[0].cells
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells[0].paragraphs[0].add_run().add_picture(str(MENOUFIA_LOGO), width=Inches(0.85))
    p = cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p)
    add_text(p, "جامعة المنوفية", size=15, bold=True)
    p.add_run().add_break()
    add_text(p, "كلية الذكاء الاصطناعي", size=15, bold=True)
    cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cells[2].paragraphs[0].add_run().add_picture(str(FACULTY_AI_LOGO), width=Inches(0.9))

    for _ in range(1):
        doc.add_paragraph()

    for text, size, bold, color in [
        ("عنوان المشروع", 13, False, None),
        (data["arabic_project_title"], 18, True, None),
        ("مشروع تخرج مقدم", 13, False, None),
        ("للحصول على درجة البكالوريوس في الذكاء الاصطناعي", 13, False, None),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_rtl(p)
        add_text(p, text, size=size, bold=bold, color=color)
        if text == data["arabic_project_title"]:
            pass

    for _ in range(1):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p)
    add_text(p, "فريق العمل:", size=15, bold=True)
    add_team_table(doc, data["arabic_team"], rtl=True)

    for _ in range(1):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p)
    add_text(p, "الإشراف:", size=15, bold=True)
    for line in data["arabic_supervisor_lines"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_rtl(p)
        add_text(p, line, size=13, bold=True if line.startswith("ا.د") else False,
                 color=RGBColor(255, 0, 0) if line.startswith("ا.د") else None)

    for _ in range(1):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl(p)
    add_text(p, data["arabic_term"], size=13)


# -----------------------------------------------------------------------------
# Content loading
# -----------------------------------------------------------------------------

IMAGE_RE = re.compile(r"^!\[(?P<caption>.*?)]\((?P<path>.*?)\)(?:\{width=(?P<width>[0-9.]+)\})?$")
NUMBERED_RE = re.compile(r"^\d+\.\s+(?P<text>.+)$")
HEADING_RE = re.compile(r"^(?P<marks>#{2,6})\s+(?P<text>.+)$")


def parse_markdown_table_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def parse_chapter_markdown(path: Path) -> list[dict]:
    """Parse the small Markdown subset used by chapter content files."""
    if not path.exists():
        raise FileNotFoundError(f"Chapter Markdown file not found: {path}")

    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[dict] = []
    paragraph_lines: list[str] = []
    i = 0

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append({"type": "paragraph", "text": " ".join(paragraph_lines)})
            paragraph_lines.clear()

    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped == PAGE_BREAK_MARKER:
            flush_paragraph()
            blocks.append({"type": "page_break"})
            i += 1
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match:
            flush_paragraph()
            blocks.append({
                "type": "heading",
                "text": heading_match.group("text"),
                "level": len(heading_match.group("marks")),
            })
            i += 1
            continue

        image_match = IMAGE_RE.match(stripped)
        if image_match:
            flush_paragraph()
            block = {
                "type": "figure",
                "path": image_match.group("path"),
                "caption": image_match.group("caption"),
            }
            if image_match.group("width"):
                block["width"] = float(image_match.group("width"))
            blocks.append(block)
            i += 1
            continue

        numbered_match = NUMBERED_RE.match(stripped)
        if numbered_match:
            flush_paragraph()
            items = []
            while i < len(lines):
                item_match = NUMBERED_RE.match(lines[i].strip())
                if not item_match:
                    break
                items.append(item_match.group("text"))
                i += 1
            blocks.append({"type": "list", "items": items})
            continue

        if stripped.startswith("Table: "):
            flush_paragraph()
            caption = stripped.removeprefix("Table: ").strip()
            i += 1
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i + 1 >= len(lines) or not lines[i].strip().startswith("|"):
                raise ValueError(f"Table caption without Markdown table in {path}: {caption}")
            headers = parse_markdown_table_row(lines[i])
            i += 2  # Skip header and separator rows.
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_markdown_table_row(lines[i]))
                i += 1
            blocks.append({
                "type": "table",
                "caption": caption,
                "headers": headers,
                "rows": rows,
            })
            continue

        paragraph_lines.append(stripped)
        i += 1

    flush_paragraph()
    return blocks


def expand_chapter_files(data: dict, *, content_dir: Path) -> dict:
    for chapter in data.get("chapters", []):
        if "file" in chapter:
            chapter["blocks"] = parse_chapter_markdown(content_dir / chapter["file"])
    return data

def load_book_data(path: Path = BOOK_JSON) -> dict:
    """Load structured book content from JSON."""
    if not path.exists():
        raise FileNotFoundError(f"Book content JSON not found: {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    return expand_chapter_files(data, content_dir=path.parent)


BOOK_DATA = load_book_data()


def build_book(data: dict, out_path: Path = OUTPUT_DOCX):
    doc = Document()
    configure_document(doc)

    add_cover_page(doc, data)
    add_front_matter(doc, data)
    for chapter in data["chapters"]:
        add_chapter(doc, chapter)
    add_bibliography(doc, data["bibliography"])
    add_arabic_summary(doc, data)
    add_arabic_cover_page(doc, data)

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    output = build_book(BOOK_DATA)
    print(f"Created: {output}")
